"""Structure-retention metric: the reference is independent of the extractors, and each failure mode lowers the right
component (found / locator / order). Synthetic documents only."""

from __future__ import annotations

import ast
import io
from pathlib import Path

from pii_shield.eval.structure import reference_units, score
from pii_shield.extraction import extract
from pii_shield.models import Locator
from pii_shield.settings import Settings

S = Settings(use_ner=False, ocr_workers=1)


def _docx() -> bytes:
    import docx

    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Internal training material"
    d.add_paragraph("Purpose of the onboarding guide")
    d.add_paragraph("Scope covers vendor intake and renewal")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Group", "Members"
    t.cell(1, 0).text, t.cell(1, 1).text = "Legal review", "Alder Quint"
    d.add_paragraph("Closing remarks for the reader")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6), Inches(1))
    tb.text_frame.text = "Organisation overview"
    tb.text_frame.add_paragraph().text = "Second paragraph of the box"
    tbl = s.shapes.add_table(2, 2, Inches(0.5), Inches(2), Inches(6), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "Role", "Holder"
    tbl.cell(1, 0).text, tbl.cell(1, 1).text = "Chair", "Alder Quint"
    s.notes_slide.notes_text_frame.text = "Speaker reminder for this slide"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _pdf() -> bytes:
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((50, 72), "Page one carries the policy text", fontsize=11)
    doc.new_page()  # blank: not a unit
    out = doc.tobytes()
    doc.close()
    return out


def test_docx_fully_retained() -> None:
    data = _docx()
    units = reference_units(data, "s.docx")
    kinds = {u.kind for u in units}
    assert {"paragraph", "table-cell", "header"} <= kinds
    r = score(units, extract(data, "s.docx", S))
    assert r["retention"] == 1.0, r


def test_pptx_fully_retained() -> None:
    data = _pptx()
    units = reference_units(data, "s.pptx")
    assert {"slide-text", "slide-table", "notes"} <= {u.kind for u in units}
    r = score(units, extract(data, "s.pptx", S))
    assert r["retention"] == 1.0, r


def test_pdf_units_are_inked_pages() -> None:
    data = _pdf()
    units = reference_units(data, "s.pdf")
    assert [u.index for u in units] == [1]
    assert score(units, extract(data, "s.pdf", S))["retention"] == 1.0


def test_dropped_block_counts_as_not_found() -> None:
    data = _docx()
    units = reference_units(data, "s.docx")
    doc = extract(data, "s.docx", S)
    doc.blocks = [b for b in doc.blocks if "Scope covers" not in b.text]
    r = score(units, doc)
    assert r["found"] == r["units"] - 1 and r["retained"] == r["units"] - 1


def test_wrong_locator_is_not_retained() -> None:
    data = _docx()
    units = reference_units(data, "s.docx")
    doc = extract(data, "s.docx", S)
    for b in doc.blocks:
        if b.text == "Alder Quint":
            b.locator = Locator(file=b.locator.file, unit="body", part="", block_id=b.id)  # cell text claimed as body
    r = score(units, doc)
    assert r["found"] == r["units"] and r["locator_ok"] == r["units"] - 1 and r["retained"] == r["units"] - 1


def test_out_of_order_block_is_not_retained() -> None:
    data = _docx()
    units = reference_units(data, "s.docx")
    doc = extract(data, "s.docx", S)
    i = next(k for k, b in enumerate(doc.blocks) if b.text.startswith("Closing remarks"))
    doc.blocks.insert(0, doc.blocks.pop(i))  # last paragraph moved to the front
    r = score(units, doc)
    assert r["found"] == r["units"] and r["in_order"] == r["units"] - 1


def test_reference_does_not_reuse_extractors() -> None:
    src = Path(__file__).parents[1] / "src" / "pii_shield" / "eval" / "structure.py"
    mods = set()
    for node in ast.walk(ast.parse(src.read_text(encoding="utf-8"))):
        if isinstance(node, ast.ImportFrom):
            mods.add(node.module or "")
        elif isinstance(node, ast.Import):
            mods.update(a.name for a in node.names)
    assert not any("extraction" in m or m in ("docx", "pptx") for m in mods), mods
