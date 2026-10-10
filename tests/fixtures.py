"""In-memory synthetic documents. Nothing here is real personal data."""

from __future__ import annotations

import io
import zipfile

SYN_TEXT = [
    "Prepared by Ottilie Brennagh (EMP-40917), Risk Analytics.",
    "Reach her at o.brennagh@example.test or +1 (212) 555-0199.",
    "Customer SSN 900-55-1234, DOB 14 July 1991, passport TK4820913.",
    "Mailing address: 1180 Ashwood Lane, Apartment 4C, Jersey City, NJ 07302.",
]
# values that must never appear in sanitized output, logs or reports
SYN_VALUES = ["Ottilie Brennagh", "EMP-40917", "o.brennagh@example.test", "555-0199", "900-55-1234", "14 July 1991",
              "TK4820913", "1180 Ashwood Lane"]


def docx_bytes(paragraphs: list[str] = SYN_TEXT) -> bytes:
    import docx

    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    d.core_properties.author = ""
    d.core_properties.last_modified_by = ""
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def pptx_bytes(lines: list[str] = SYN_TEXT) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(5))
    tb.text_frame.text = lines[0]
    for line in lines[1:]:
        tb.text_frame.add_paragraph().text = line
    prs.core_properties.author = ""
    prs.core_properties.last_modified_by = ""
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def pdf_bytes(lines: list[str] = SYN_TEXT, encrypt: bool = False) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    y = 72
    for line in lines:
        page.insert_text((50, y), line, fontsize=10)
        y += 18
    kw = {}
    if encrypt:
        kw = {"encryption": pymupdf.PDF_ENCRYPT_AES_256, "user_pw": "synthetic", "owner_pw": "synthetic-owner"}
    out = doc.tobytes(**kw)
    doc.close()
    return out


def rezip(data: bytes, add: dict[str, bytes] | None = None, replace: dict[str, bytes] | None = None,
          encrypted_flag: str | None = None) -> bytes:
    """Copy an OOXML zip, optionally adding/replacing entries or flagging one entry as encrypted."""
    add, replace = add or {}, replace or {}
    src = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = replace.get(info.filename, src.read(info.filename))
            zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            zi.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(zi, body)
        for name, body in add.items():
            dst.writestr(zipfile.ZipInfo(name), body, compress_type=zipfile.ZIP_DEFLATED)
    out = bytearray(buf.getvalue())
    if encrypted_flag:  # zipfile resets flag_bits on write, so set "encrypted" (bit 0) in both headers by hand
        info = zipfile.ZipFile(io.BytesIO(bytes(out))).getinfo(encrypted_flag)
        out[info.header_offset + 6] |= 0x1  # local file header: general purpose flag at +6
        cd = out.find(b"PK\x01\x02" + bytes(out[info.header_offset + 4:info.header_offset + 6]))
        while cd != -1:
            n = int.from_bytes(out[cd + 28:cd + 30], "little")
            if out[cd + 46:cd + 46 + n] == encrypted_flag.encode():
                out[cd + 8] |= 0x1  # central directory: general purpose flag at +8
                break
            cd = out.find(b"PK\x01\x02", cd + 4)
    return bytes(out)
