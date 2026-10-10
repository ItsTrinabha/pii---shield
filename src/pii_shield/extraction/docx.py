"""DOCX extraction in document order: paragraphs, nested tables, text boxes, headers/footers, footnotes/endnotes,
comments, tracked deletions, hidden text, alt text, embedded images (OCR), metadata, then an independent XML sweep."""

from __future__ import annotations

import io
import zipfile

import docx
from lxml import etree

from ..errors import ExtractionError
from ..models import Block, DocStats, ExtractedDocument, Locator
from ..security import safe_name, sha256_hex, validate_zip_package
from ..settings import Settings
from .ooxml import (
    MC, NS, W, hyperlink_blocks, image_blocks, metadata_blocks, norm, ocr_images, para_text, parse_xml, read_part, sweep,
)


def extract_docx(data: bytes, file_name: str, settings: Settings) -> ExtractedDocument:
    name = safe_name(file_name)
    zf = validate_zip_package(data, "docx", settings.limits)
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise ExtractionError("DOCX could not be parsed", "CORRUPT_FILE") from e

    blocks: list[Block] = []
    stats = DocStats()
    warnings: list[str] = []
    order = 0
    images: dict[str, bytes] = {}
    pending_images: dict[int, list[str]] = {}  # index in `blocks` after which OCR text is inserted -> image keys
    page_marker = {"n": 0}

    def add(text: str, unit: str, part: str = "", source: str = "native") -> int:
        nonlocal order
        if not text.strip():
            return -1
        loc = Locator(file=name, unit=unit, index=None, part=part, block_id=f"{unit[:3]}-{order}")
        blocks.append(Block(id=loc.block_id, text=text, locator=loc, source=source, order=order))
        order += 1
        return len(blocks) - 1

    def rel_image(owner_part, rid: str) -> tuple[str, bytes] | None:
        try:
            part = owner_part.related_parts[rid]
        except KeyError:
            return None
        pn = str(part.partname).lstrip("/")
        if not pn.startswith("word/media/") and "media" not in pn:
            return None
        return pn, part.blob

    def handle_paragraph(p: etree._Element, unit: str, part_prefix: str, owner_part) -> None:
        text = para_text(p, True)
        idx = add(text, unit, part_prefix)
        # tracked deletions / hidden text / alt text / field codes
        deleted = "".join((e.text or "") for e in p.iter(W + "delText"))
        if deleted.strip():
            add(deleted, unit, f"{part_prefix} tracked-deletion".strip())
            stats.hidden_content_blocks += 1
        if any(True for _ in p.iter(W + "vanish")):
            stats.hidden_content_blocks += 1
        for dp in p.iter("{%s}docPr" % NS["wp"]):
            for k in ("descr", "title"):
                v = dp.get(k)
                if v and v.strip():
                    add(v, unit, f"{part_prefix} alt-text".strip())
        for it in p.iter(W + "instrText"):
            if it.text and ("@" in it.text or "HYPERLINK" in it.text.upper()):
                add(it.text.strip(), unit, f"{part_prefix} field-code".strip())
        # text boxes: their paragraphs are blocks of their own
        for tb in p.iterfind(".//" + W + "txbxContent"):
            anc = tb.getparent()
            fallback = False
            while anc is not None:
                if anc.tag == MC + "Fallback":
                    fallback = True
                    break
                anc = anc.getparent()
            if fallback:
                continue
            walk_container(tb, unit, f"{part_prefix} textbox".strip(), owner_part)
        # images: record for OCR right after this paragraph
        rids = [b.get("{%s}embed" % NS["r"]) for b in p.iter("{%s}blip" % NS["a"])]
        rids += [b.get("{%s}id" % NS["r"]) for b in p.iter("{%s}imagedata" % NS["v"])]
        for rid in [r for r in rids if r]:
            got = rel_image(owner_part, rid)
            if got:
                key, blob = got
                images.setdefault(key, blob)
                pending_images.setdefault(len(blocks) - 1, []).append(key)

    def walk_container(parent: etree._Element, unit: str, part_prefix: str, owner_part) -> None:
        for ch in parent:
            if ch.tag == W + "p":
                handle_paragraph(ch, unit, part_prefix, owner_part)
            elif ch.tag == W + "tbl":
                for ri, tr in enumerate(ch.iterfind(W + "tr")):
                    for ci, tc in enumerate(tr.iterfind(W + "tc")):
                        walk_container(tc, unit, f"{part_prefix} table r{ri + 1}c{ci + 1}".strip(), owner_part)
            elif ch.tag in (W + "sdt",):
                content = ch.find(W + "sdtContent")
                if content is not None:
                    walk_container(content, unit, part_prefix, owner_part)
            elif ch.tag in (W + "ins", W + "moveTo", W + "customXml", W + "smartTag"):
                walk_container(ch, unit, part_prefix, owner_part)

    # ---- body ---------------------------------------------------------------------------------------------
    body = d.element.body
    walk_container(body, "body", "", d.part)
    stats.page_count = 0  # DOCX has no fixed pagination without rendering

    # ---- headers / footers (deduplicated by part) -------------------------------------------------------------
    seen_parts: set[str] = set()
    for si, sec in enumerate(d.sections, 1):
        for kind, hf in (("header", sec.header), ("first-page header", sec.first_page_header), ("even-page header", sec.even_page_header),
                         ("footer", sec.footer), ("first-page footer", sec.first_page_footer), ("even-page footer", sec.even_page_footer)):
            try:
                if hf.is_linked_to_previous:
                    continue
                pn = str(hf.part.partname)
                if pn in seen_parts:
                    continue
                seen_parts.add(pn)
                unit = "header" if "header" in kind else "footer"
                walk_container(hf._element, unit, f"section {si} {kind}", hf.part)
            except Exception:  # noqa: BLE001
                warnings.append(f"section {si}: {kind} could not be read (completeness sweep covers it)")

    # ---- footnotes, endnotes, comments ----------------------------------------------------------------------
    names = set(zf.namelist())
    for part, tag, unit in (("word/footnotes.xml", "footnote", "footnote"), ("word/endnotes.xml", "endnote", "endnote"),
                            ("word/comments.xml", "comment", "comment")):
        if part not in names:
            continue
        root = parse_xml(read_part(zf, part, settings), settings.limits.max_xml_part_bytes)
        for el in root.iter(W + tag):
            if el.get(W + "type") in ("separator", "continuationSeparator"):
                continue
            if unit == "comment":
                for k in ("author", "initials"):
                    v = el.get(W + k)
                    if v and v.strip():
                        add(v, unit, f"{k}")
            for p in el.iter(W + "p"):
                add(para_text(p, True), unit, "")
                stats.hidden_content_blocks += 1 if unit == "comment" else 0

    # ---- images: OCR and insert at their position -----------------------------------------------------------
    img_results, unsupported = ocr_images(images, settings) if images else ({}, [])
    stats.unsupported_media.extend(f"{k}: unsupported image format, not analysed" for k in unsupported)
    # SVG images are XML text: read their <text> content directly
    for k in unsupported:
        if k.lower().endswith(".svg"):
            try:
                root = parse_xml(images[k], settings.limits.max_xml_part_bytes)
                txt = " ".join(t.strip() for t in root.itertext() if t and t.strip())
                if txt:
                    add(txt, "body", f"image:{k}", "native")
                    stats.unsupported_media = [u for u in stats.unsupported_media if not u.startswith(k)]
            except Exception:  # noqa: BLE001
                pass
    if pending_images:
        rebuilt: list[Block] = []
        placed: set[str] = set()
        for i, b in enumerate(blocks):
            rebuilt.append(b)
            for key in pending_images.get(i, []):
                pg = img_results.get(key)
                if pg is None or key in placed:
                    continue
                placed.add(key)
                for ib in image_blocks(name, b.locator.unit, None, key, pg, 0):
                    rebuilt.append(ib)
                stats.ocr_blocks += len(pg.paragraphs)
        # images referenced before any block exists (index -1)
        for key in pending_images.get(-1, []):
            pg = img_results.get(key)
            if pg is not None and key not in placed:
                placed.add(key)
                for ib in image_blocks(name, "body", None, key, pg, 0):
                    rebuilt.insert(0, ib)
        blocks = rebuilt
        for i, b in enumerate(blocks):
            b.order = i
        order = len(blocks)

    # ---- metadata, link targets, independent sweep ------------------------------------------------------------
    for b in metadata_blocks(zf, name, settings, order):
        blocks.append(b)
        order = b.order + 1
    for b in hyperlink_blocks(zf, name, settings, order):
        blocks.append(b)
        order = b.order + 1
    corpus = norm(" \n ".join(b.text for b in blocks))
    sw = sweep(zf, "docx", name, corpus, settings, order)
    blocks.extend(sw)
    stats.sweep_blocks = len(sw)
    if sw:
        warnings.append(f"{len(sw)} text unit(s) were found only by the XML completeness sweep (hidden/unsupported structures)")

    if not any(b.locator.unit == "body" for b in blocks):
        warnings.append("no body text could be extracted")

    _c = [b.ocr_confidence for b in blocks if b.source == "ocr" and b.ocr_confidence is not None]
    stats.mean_ocr_confidence = round(sum(_c) / len(_c), 1) if _c else None
    stats.block_count = len(blocks)
    stats.ocr_chars = sum(len(b.text) for b in blocks if b.source == "ocr")
    stats.native_chars = sum(len(b.text) for b in blocks if b.source != "ocr")
    return ExtractedDocument(file_name=name, file_type="docx", sha256=sha256_hex(data), size_bytes=len(data),
                             blocks=blocks, stats=stats, warnings=warnings, source_bytes=data)

